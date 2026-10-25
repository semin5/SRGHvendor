"""Import the supplied vendor workbook into the new srghvendor database.

Only run against the writable MySQL primary. MYSQL_PWD must be supplied through
the process environment; no credentials or source records are logged.
"""

import argparse
import os
import re
import subprocess
from pathlib import Path

import openpyxl


TABLES = {
    "업무사이트": (
        "work_sites",
        1,
        ("업무명칭", "사이트", "ID", "PW", "업체", "담당자", "연락처", "비고"),
        (1, 2, 3, 4, 5, 6, 7, 8),
        ("work_name", "site_url", "login_id", "password", "vendor_name",
         "contact_person", "contact_phone", "remarks"),
    ),
    "업체 정보": (
        "vendor_info",
        1,
        ("업체명", "제품", "담당자", "전화번호", "사무실 전화", "주소", "홈페이지", "택배이용"),
        (1, 2, 3, 4, 5, 6, 7, 8),
        ("vendor_name", "product", "contact_person", "phone_number",
         "office_phone", "address", "website", "courier_service"),
    ),
    "유지보수업체 리스트": (
        "maintenance_vendors",
        4,
        ("유지보수업체", "전화번호"),
        (2, 3),
        ("vendor_name", "phone_number"),
    ),
}


def mysql(args, sql):
    result = subprocess.run(
        [args.mysql, "--default-character-set=utf8mb4", "--batch", "--skip-column-names",
         "-h", args.host, "-P", str(args.port), "-u", args.user],
        input=sql.encode("utf-8"),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(f"MySQL command failed (exit {result.returncode}); source values were not printed")
    return result.stdout.decode("utf-8").strip()


def sql_value(value):
    if value is None:
        return "NULL"
    if isinstance(value, int):
        return str(value)
    return f"CONVERT(0x{str(value).encode('utf-8').hex()} USING utf8mb4)"


def sheet_rows(sheet, header_row, headers, columns, fields):
    actual = tuple(str(sheet.cell(header_row, column).value or "").strip() for column in columns)
    if actual != headers:
        raise ValueError(f"Unexpected headers in {sheet.title}")
    result = []
    for excel_row, row in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
        if not any(cell.value is not None for cell in row):
            continue
        values = [row[column - 1].value for column in columns]
        if any(cell.value is not None for i, cell in enumerate(row, 1) if i not in columns):
            raise ValueError(f"Unexpected extra data in {sheet.title} row {excel_row}")
        result.append((excel_row, fields, values))
    return result


def video_rows(sheet):
    year_label = str(sheet["A1"].value or "")
    match = re.fullmatch(r"(\d{4})년", year_label)
    if not match or str(sheet["B1"].value or "").strip() != "변경 내용":
        raise ValueError("Unexpected video inspection header")
    year = int(match.group(1))
    result = []
    for excel_row, row in enumerate(sheet.iter_rows(min_row=2), 2):
        if not any(cell.value is not None for cell in row):
            continue
        month_label = str(row[0].value or "")
        month_match = re.fullmatch(r"(\d{1,2})월", month_label)
        if not month_match or not 1 <= int(month_match.group(1)) <= 12:
            raise ValueError(f"Unexpected video inspection month in row {excel_row}")
        result.append((excel_row,
                       ("inspection_year", "inspection_month", "change_details"),
                       [year, int(month_match.group(1)), row[1].value]))
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--mysql", required=True)
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--host", default="192.168.1.77")
    parser.add_argument("--port", type=int, default=3307)
    parser.add_argument("--user", default="semin5")
    args = parser.parse_args()
    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    if set(workbook.sheetnames) != set(TABLES) | {"영상정보월간점검일지"}:
        raise ValueError("Workbook sheet list differs from the expected four sheets")
    batches = {}
    for title, (table, header_row, headers, columns, fields) in TABLES.items():
        batches[table] = sheet_rows(workbook[title], header_row, headers, columns, fields)
    batches["video_monthly_inspections"] = video_rows(workbook["영상정보월간점검일지"])
    expected = {table: len(rows) for table, rows in batches.items()}
    if args.validate_only:
        for table, count in expected.items():
            print(f"{table}: {count} source rows validated")
        return
    if not os.environ.get("MYSQL_PWD"):
        raise RuntimeError("MYSQL_PWD must be set")
    if mysql(args, "SELECT @@read_only;") != "0":
        raise RuntimeError("Target is not the writable primary")
    if mysql(args, "SHOW DATABASES LIKE 'srghvendor';"):
        raise RuntimeError("srghvendor already exists; refusing to overwrite it")

    schema = Path(__file__).resolve().parents[1] / "db" / "schema.sql"
    mysql(args, schema.read_text(encoding="utf-8"))
    statements = ["USE srghvendor;", "START TRANSACTION;"]
    for table, rows in batches.items():
        for source_row, fields, values in rows:
            columns = ", ".join(("source_row", *fields))
            literals = ", ".join(sql_value(value) for value in (source_row, *values))
            statements.append(f"INSERT INTO {table} ({columns}) VALUES ({literals});")
    statements.append("COMMIT;")
    mysql(args, "\n".join(statements))

    for table, count in expected.items():
        actual = int(mysql(args, f"SELECT COUNT(*) FROM srghvendor.{table};"))
        if actual != count:
            raise RuntimeError(f"{table}: expected {count} rows, found {actual}")
        print(f"{table}: {actual} rows imported")


if __name__ == "__main__":
    main()
